"""결과 저장: tidy CSV · 지표별 피벗 CSV · 엑셀 · 커버리지 리포트."""

from __future__ import annotations

import csv
import logging
from collections import Counter, defaultdict
from pathlib import Path

from .assemble import KOREAN_HEADERS, RECORD_COLUMNS

log = logging.getLogger(__name__)

# 지표별 피벗 대상: (파일명 접미사, 레코드 컬럼, 시트/제목)
PIVOT_METRICS = [
    ("total_assets", "total_assets", "자산총액"),
    ("revenue_quarter", "revenue_quarter", "매출액(분기)"),
    ("revenue_cumulative", "revenue_cumulative", "매출액(누적)"),
    ("operating_income_quarter", "operating_income_quarter", "영업이익(분기)"),
    ("operating_income_cumulative", "operating_income_cumulative", "영업이익(누적)"),
    ("employees", "employees", "인력수"),
]

ID_COLUMNS = ["corp_code", "corp_name", "stock_code", "market", "induty_code"]


def write_tidy_csv(rows: list[dict], path: Path, korean_header: bool = True) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        header = [KOREAN_HEADERS[c] for c in RECORD_COLUMNS] if korean_header else RECORD_COLUMNS
        writer.writerow(header)
        for row in rows:
            writer.writerow(["" if row.get(c) is None else row.get(c) for c in RECORD_COLUMNS])
    log.info("저장: %s (%d행)", path, len(rows))
    return path


def pivot(rows: list[dict], value_column: str) -> tuple[list[str], list[list]]:
    """기업 1행 × 분기 열 형태로 변환."""
    periods = sorted({r["period"] for r in rows})
    by_corp: dict[str, dict] = {}
    values: dict[str, dict[str, object]] = defaultdict(dict)
    for row in rows:
        by_corp.setdefault(row["corp_code"], row)
        values[row["corp_code"]][row["period"]] = row.get(value_column)

    header = [KOREAN_HEADERS[c] for c in ID_COLUMNS] + periods
    out = []
    for corp_code, meta in sorted(by_corp.items(), key=lambda kv: (kv[1]["market"], kv[1]["stock_code"])):
        line = [meta.get(c, "") for c in ID_COLUMNS]
        line += ["" if values[corp_code].get(p) is None else values[corp_code].get(p) for p in periods]
        out.append(line)
    return header, out


def write_pivot_csvs(rows: list[dict], out_dir: Path) -> list[Path]:
    written = []
    for suffix, column, _title in PIVOT_METRICS:
        header, body = pivot(rows, column)
        path = out_dir / f"pivot_{suffix}.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh)
            writer.writerow(header)
            writer.writerows(body)
        written.append(path)
        log.info("저장: %s (%d개사)", path, len(body))
    return written


def write_excel(rows: list[dict], path: Path) -> Path | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        log.warning("openpyxl 이 없어 엑셀 파일을 건너뜁니다 (pip install openpyxl)")
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "전체(tidy)"
    ws.append([KOREAN_HEADERS[c] for c in RECORD_COLUMNS])
    for row in rows:
        ws.append([row.get(c) for c in RECORD_COLUMNS])
    _style_sheet(ws, Font, Alignment, get_column_letter)

    for _suffix, column, title in PIVOT_METRICS:
        header, body = pivot(rows, column)
        sheet = wb.create_sheet(title[:31])
        sheet.append(header)
        for line in body:
            sheet.append([v if v != "" else None for v in line])
        _style_sheet(sheet, Font, Alignment, get_column_letter)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    log.info("저장: %s", path)
    return path


def _style_sheet(ws, Font, Alignment, get_column_letter) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    for idx, cell in enumerate(ws[1], start=1):
        width = max(10, min(22, len(str(cell.value or "")) * 2 + 2))
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, int) and abs(cell.value) >= 1000:
                cell.number_format = "#,##0"


def coverage_report(rows: list[dict], periods: list[str]) -> str:
    """분기·시장별 수집 커버리지 요약(마크다운)."""
    corps = {(r["corp_code"], r["market"]) for r in rows}
    market_totals = Counter(market for _, market in corps)

    lines = ["# 수집 커버리지", "", "## 유니버스", ""]
    for market, count in sorted(market_totals.items()):
        lines.append(f"- {market}: {count:,}개사")
    lines += ["", "## 분기별 데이터 확보율", "",
              "| 기간 | 시장 | 대상 | 자산총액 | 매출액 | 영업이익 | 인력수 |",
              "|---|---|---:|---:|---:|---:|---:|"]

    for period in periods:
        for market in sorted(market_totals):
            subset = [r for r in rows if r["period"] == period and r["market"] == market]
            if not subset:
                continue
            total = len(subset)
            def rate(col):
                got = sum(1 for r in subset if r.get(col) is not None)
                return f"{got:,} ({got / total * 100:.1f}%)"
            lines.append(
                f"| {period} | {market} | {total:,} | {rate('total_assets')} | "
                f"{rate('revenue_cumulative')} | {rate('operating_income_cumulative')} | {rate('employees')} |"
            )

    sources = Counter(r["revenue_quarter_source"] for r in rows if r.get("revenue_quarter") is not None)
    lines += ["", "## 분기 매출액 산출 근거", ""]
    for source, count in sources.most_common():
        lines.append(f"- {source}: {count:,}건")

    non_dec = {r["corp_name"] for r in rows if r.get("non_december_fy") == "Y"}
    lines += ["", "## 참고", "",
              f"- 12월 결산이 아닌 기업 {len(non_dec):,}개사: 분기 라벨(2025Q1 등)은 회사의 회계연도 기준입니다.",
              "- 인력수는 사업보고서·반기보고서에만 공시되며, 1·3분기는 직전 공시치를 이월한 참고값입니다."]
    return "\n".join(lines) + "\n"
